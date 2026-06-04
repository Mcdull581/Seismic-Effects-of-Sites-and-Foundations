# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

import json
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INPUT_JSON = os.path.join(BASE_DIR, "data", "seismic_data_full.json")
OUTPUT_DOCX = os.path.join(BASE_DIR, "output", "核心规范与全国设防参数汇总_纯表格版.docx")


def set_cell_shading(cell, color="D9E2F3"):
    """设置单元格背景色"""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text, bold=False, font_size=9):
    """设置单元格文本，居中对齐"""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(font_size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def create_word_table(doc, title, headers, rows_data, desc=""):
    """
    通用表格生成函数
    headers: 表头列表，如 ["土的类型", "岩土名称和性状", "剪切波速范围(m/s)"]
    rows_data: 二维列表，每行是与headers等长的值列表，如 [["岩石", "坚硬...", "vs>800"], ...]
    """
    h = doc.add_heading(title, level=2)
    for run in h.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    if desc:
        p = doc.add_paragraph()
        run = p.add_run(desc)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(100, 100, 100)
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    num_cols = len(headers)
    num_rows = 1 + len(rows_data)
    table = doc.add_table(rows=num_rows, cols=num_cols, style='Table Grid')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    for i, col in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, col, bold=True, font_size=10)
        set_cell_shading(cell, "1F4E79")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)

    # 数据行
    for r_idx, row in enumerate(rows_data):
        bg = "F2F7FB" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx in range(num_cols):
            cell = table.rows[r_idx + 1].cells[c_idx]
            val = row[c_idx] if c_idx < len(row) else ""
            if val is None:
                val = "—"
            set_cell_text(cell, str(val), font_size=9)
            set_cell_shading(cell, bg)

    doc.add_paragraph("")
    print(f"  [OK] {title} ({len(rows_data)} 行)")


def extract_and_generate(json_path, output_path):
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    # 文档标题
    title = doc.add_heading("核心规范与全国设防参数汇总", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("（纯表格版 — 依据 GB 50011-2010 / GB 18306-2015 / CECS 55:2017）")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(100, 100, 100)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    doc.add_paragraph("")

    # ================================================================
    # 第一部分：核心规范数据表
    # ================================================================
    doc.add_heading("第一部分：核心规范数据表", level=1)

    # ---------- 表4.1.3 ----------
    t = data["table_4_1_3_soil_type"]
    create_word_table(doc,
        "表4.1.3  土的类型与剪切波速范围",
        ["土的类型", "岩土名称和性状", "剪切波速范围(m/s)"],
        [[r["type"], r["description"], r["vs_range"]] for r in t["data"]],
        t.get("description", "")
    )

    # ---------- 表4.1.6 ----------
    t = data["table_4_1_6_site_classification"]
    create_word_table(doc,
        "表4.1.6  场地类别划分",
        ["等效剪切波速(m/s)", "I₀", "I₁", "II", "III", "IV"],
        [[r["vs_condition"], r["I0"], r["I1"], r["II"], r["III"], r["IV"]] for r in t["data"]],
        t.get("description", "")
    )

    # ---------- 表4.3.3 ----------
    t = data["table_4_3_3_liquefaction_depth"]
    create_word_table(doc,
        "表4.3.3  液化土特征深度 (m)",
        ["饱和土类别", "7度", "8度", "9度"],
        [[r["soil_type"], r["intensity_7"], r["intensity_8"], r["intensity_9"]] for r in t["data"]],
        t.get("description", "")
    )

    # ---------- 表4.3.4 N0 ----------
    t = data["table_4_3_4_N0"]
    create_word_table(doc,
        "表4.3.4  液化判别标准贯入锤击数基准值 N₀",
        ["设计基本地震加速度(g)", "N₀"],
        [[r["pga"], r["N0"]] for r in t["data"]],
        t.get("description", "")
    )

    # ---------- 表4.3.4 beta ----------
    t = data["table_4_3_4_beta"]
    create_word_table(doc,
        "表4.3.4  调整系数 β",
        ["设计地震分组", "β"],
        [[r["group"], r["beta"]] for r in t["data"]],
        t.get("description", "")
    )

    # ---------- 表4.3.5 ----------
    t = data["table_4_3_5_liquefaction_grade"]
    create_word_table(doc,
        "表4.3.5  液化等级",
        ["液化等级", "液化指数 I_lE 范围"],
        [[r["grade"], r["description"]] for r in t["data"]],
        t.get("description", "")
    )

    # ---------- 表4.3.6 ----------
    t = data["table_4_3_6_liquefaction_measures"]
    create_word_table(doc,
        "表4.3.6  抗液化措施",
        ["抗震设防类别", "轻微液化", "中等液化", "严重液化"],
        [[r["category"], r["slight"], r["moderate"], r["severe"]] for r in t["data"]],
        t.get("description", "")
    )

    # ---------- 第4.1.1条 抗震地段划分 ----------
    t = data["table_4_1_1_site_classification_type"]
    # 这个表的data结构特殊，conditions是列表
    for item in t["data"]:
        conds = item["conditions"]
        # 每个地段类别单独做一张小表
        create_word_table(doc,
            f"第4.1.1条  {item['category']}",
            ["判定条件"],
            [[c] for c in conds],
        )

    # ---------- 软土震陷初判阈值 ----------
    ss = data["soft_soil_subsidenc"]
    wvt = ss["initial_judgment"]["wave_velocity_threshold"]
    create_word_table(doc,
        "软弱土震陷 — 初判阈值表",
        ["设防烈度", "等效剪切波速阈值(m/s)", "承载力特征值阈值(kPa)"],
        [[r["intensity"], r["vse_threshold"], r["fak_threshold"]] for r in wvt["data"]],
        wvt.get("description", "")
    )

    # ---------- 软土震陷估算值 ----------
    se = ss["subsidence_estimate"]
    create_word_table(doc,
        "软弱土震陷 — 震陷估算值表",
        ["设防烈度", "Vse=90~140 m/s", "Vse=140~200 m/s"],
        [[r["intensity"], r["vse_90_140"], r["vse_140_200"]] for r in se["data"]],
        se.get("description", "")
    )

    # ---------- 液化初判条件 ----------
    liq = data["liquefaction_initial_judgment"]
    create_word_table(doc,
        "砂土液化 — 初判条件汇总",
        ["条款", "判定逻辑"],
        [
            [liq["condition_1"]["clause"], liq["condition_1"]["logic"]],
            [liq["condition_2"]["clause"], liq["condition_2"]["logic"]],
            [liq["condition_3"]["clause"],
             "du > d0+db-2  或  dw > d0+db-3  或  du+dw > 1.5d0+2db-4.5"],
        ]
    )

    # ---------- 公式汇总 ----------
    formulas = data["formulas"]
    create_word_table(doc,
        "核心公式汇总",
        ["公式名称", "公式表达式"],
        [
            ["等效剪切波速", f"{formulas['equivalent_wave_velocity']['eq1']}；  {formulas['equivalent_wave_velocity']['eq2']}"],
            ["液化判别临界值", formulas["liquefaction_critical_N"]["formula"]],
            ["液化指数", formulas["liquefaction_index"]["formula"]],
            ["权函数", formulas["weight_function"]["formula"]],
        ]
    )

    # ================================================================
    # 第二部分：全国抗震设防烈度表（附录A）
    # ================================================================
    doc.add_heading("第二部分：全国抗震设防烈度表（附录A · GB 50011-2010）", level=1)

    aa_data = data["appendix_a_seismic_intensity"]["data"]
    provinces = {}
    for rec in aa_data:
        prov = rec["province"]
        if prov not in provinces:
            provinces[prov] = []
        provinces[prov].append(rec)

    print(f"\n  附录A: {len(aa_data)} 条记录，{len(provinces)} 个省级行政区")

    for prov_name in sorted(provinces.keys()):
        recs = provinces[prov_name]
        create_word_table(doc,
            f"{prov_name}",
            ["城镇", "设防烈度(度)", "设计基本地震加速度(g)", "设计地震分组"],
            [[r["city"], r["intensity"], r["pga"], r["group"]] for r in recs]
        )

    # ================================================================
    # 第三部分：全国地震动参数表（附录C）
    # ================================================================
    doc.add_heading("第三部分：全国地震动参数表（附录C · GB 18306-2015）", level=1)

    ac_data = data["appendix_c_ground_motion"]["data"]
    provinces_c = {}
    for rec in ac_data:
        prov = rec["province"]
        if prov not in provinces_c:
            provinces_c[prov] = []
        provinces_c[prov].append(rec)

    for prov_name in sorted(provinces_c.keys()):
        recs = provinces_c[prov_name]
        create_word_table(doc,
            f"{prov_name}",
            ["区域", "基本地震动峰值加速度(g)", "特征周期分区", "特征周期值(s)"],
            [[r["area"], r["pga"], r["tg_group"], r["tg_value"]] for r in recs]
        )

    # 保存
    doc.save(output_path)
    print(f"\n  文档已保存: {output_path}")


if __name__ == "__main__":
    print("=" * 50)
    print("  纯表格版 Word 文档提取脚本")
    print("=" * 50)
    extract_and_generate(INPUT_JSON, OUTPUT_DOCX)
    print("\n  纯表格版 Word 文档已生成，请查阅。")
